"""
Lambda: S3-triggered (ObjectCreated on raw/*)

Reads a raw uploaded file, asks Claude (via Bedrock converse) to turn it
into a structured training module matching the schema already used in
the production pipeline's lesson JSON (title + content_blocks), and writes
the result to processed/<name>.json for canvasAITrainer-preprod to index.

Current scope: real end-to-end extraction for .txt, .pdf, and .pptx.
.docx/.doc/.ppt/.mp4/.mov are logged and skipped rather than faked — see
backend/preprod/README.md for what real extraction would require for each
remaining format.
"""

import json
import logging
import os
import urllib.parse
from datetime import datetime, timezone
from io import BytesIO

import boto3
import pdfplumber
from pptx import Presentation

logger = logging.getLogger()
logger.setLevel(logging.INFO)

GENERATION_MODEL_ID = os.environ.get(
    "GENERATION_MODEL_ID", "us.anthropic.claude-sonnet-4-5-20250929-v1:0"
)

KNOWN_UNSUPPORTED_EXTENSIONS = {".docx", ".doc", ".ppt", ".mp4", ".mov"}


def _extract_txt(body: bytes) -> str:
    return body.decode("utf-8", errors="replace")


def _extract_pdf(body: bytes) -> str:
    pages = []
    with pdfplumber.open(BytesIO(body)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
    return "\n\n".join(pages)


def _extract_pptx(body: bytes) -> str:
    prs = Presentation(BytesIO(body))
    slides_text = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = [f"[Slide {i}]"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs)
                    if text.strip():
                        lines.append(text)
        slides_text.append("\n".join(lines))
    return "\n\n".join(slides_text)


EXTRACTORS = {
    ".txt": _extract_txt,
    ".pdf": _extract_pdf,
    ".pptx": _extract_pptx,
}

MODULE_GENERATION_SYSTEM_PROMPT = (
    "You convert raw training material text into a structured training "
    "module. Respond with ONLY valid JSON (no markdown fences, no prose "
    "before or after) matching exactly this shape: "
    '{"title": "<short descriptive title>", "content_blocks": '
    '[{"tag": "p"|"li"|"h3", "text": "<block text>"}, ...]}. '
    "Break the content into logical paragraphs (tag p), bullet points for "
    "lists of steps or options (tag li), and section headers where useful "
    "(tag h3). Preserve the actual instructional content faithfully — do "
    "not invent steps that aren't in the source text."
)

s3 = boto3.client("s3")
bedrock_runtime = boto3.client("bedrock-runtime")


# Keeps generation time bounded for large source files (long PDFs/decks) —
# a PPTX without this cap ran past a 60s Lambda timeout with no error logged,
# just a silent kill, since Bedrock had no output limit and kept generating.
MAX_INPUT_CHARS = 16_000
MAX_OUTPUT_TOKENS = 8192


def _generate_module_json(raw_text: str, source_filename: str) -> dict:
    if len(raw_text) > MAX_INPUT_CHARS:
        raw_text = raw_text[:MAX_INPUT_CHARS]
        logger.warning(
            "Truncated input for %s to %d chars before sending to the model.",
            source_filename,
            MAX_INPUT_CHARS,
        )

    user_message = f"Source file: {source_filename}\n\nRaw content:\n{raw_text}"

    response = bedrock_runtime.converse(
        modelId=GENERATION_MODEL_ID,
        system=[{"text": MODULE_GENERATION_SYSTEM_PROMPT}],
        messages=[{"role": "user", "content": [{"text": user_message}]}],
        inferenceConfig={"maxTokens": MAX_OUTPUT_TOKENS},
    )
    raw_output = response["output"]["message"]["content"][0]["text"].strip()

    # Defensive: strip markdown code fences if the model adds them anyway.
    if raw_output.startswith("```"):
        raw_output = raw_output.strip("`")
        if raw_output.startswith("json"):
            raw_output = raw_output[4:].strip()

    module = json.loads(raw_output)  # let this raise if the model didn't return valid JSON
    module["source_file"] = source_filename
    module["generated_at"] = datetime.now(timezone.utc).isoformat()
    return module


def lambda_handler(event, context):
    for record in event.get("Records", []):
        bucket = record["s3"]["bucket"]["name"]
        key = urllib.parse.unquote_plus(record["s3"]["object"]["key"])

        filename = os.path.basename(key)
        ext = os.path.splitext(filename)[1].lower()

        if ext in KNOWN_UNSUPPORTED_EXTENSIONS:
            logger.warning(
                "Skipping %s: %s extraction is not implemented yet (see "
                "backend/preprod/README.md).",
                key,
                ext,
            )
            continue

        extractor = EXTRACTORS.get(ext)
        if extractor is None:
            logger.warning("Skipping %s: unrecognized extension %s.", key, ext)
            continue

        try:
            obj = s3.get_object(Bucket=bucket, Key=key)
            body = obj["Body"].read()
            raw_text = extractor(body)

            if not raw_text.strip():
                logger.warning("Skipping %s: no extractable text found.", key)
                continue

            module = _generate_module_json(raw_text, filename)

            processed_key = f"processed/{os.path.splitext(filename)[0]}.json"
            s3.put_object(
                Bucket=bucket,
                Key=processed_key,
                Body=json.dumps(module, indent=2).encode("utf-8"),
                ContentType="application/json",
            )
            logger.info("Generated module %s from %s", processed_key, key)
        except Exception:
            logger.exception("Failed to generate module for %s", key)
            # Don't re-raise: one bad file shouldn't fail the whole batch
            # of S3 event records in this invocation.
            continue
